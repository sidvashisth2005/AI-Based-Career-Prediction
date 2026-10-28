"""Extracts text and a structured profile (skills, education, experience, projects) from resumes."""

import io
import re
from datetime import date

from app.skill_extractor import extract_skills

SECTION_HEADINGS = {
    "summary": ["summary", "profile", "objective", "about me", "career objective", "professional summary"],
    "education": ["education", "academic", "academics", "qualification", "qualifications", "academic background"],
    "experience": ["experience", "work experience", "professional experience", "employment", "internship", "internships", "work history"],
    "projects": ["projects", "project", "academic projects", "personal projects", "key projects"],
    "skills": ["skills", "technical skills", "core competencies", "technologies", "tech stack", "skill set", "tools"],
    "certifications": ["certifications", "certification", "certificates", "courses", "licenses"],
    "achievements": ["achievements", "awards", "accomplishments", "honors", "extracurricular", "activities", "positions of responsibility"],
}

EDUCATION_LEVELS = [
    ("PhD", r"\b(ph\.?\s?d|doctorate)\b"),
    ("Masters", r"\b(m\.?\s?tech|m\.?\s?e\b|m\.?\s?sc|mca|mba|m\.?\s?s\b|master'?s?|pgdm|m\.?\s?com)\b"),
    ("Bachelors", r"\b(b\.?\s?tech|b\.?\s?e\b|b\.?\s?sc|bca|bba|b\.?\s?com|b\.?\s?a\b|bachelor'?s?|b\.?\s?des|undergraduate)\b"),
    ("Diploma", r"\b(diploma|polytechnic)\b"),
    ("Higher Secondary", r"\b(12th|xii|higher secondary|hsc|intermediate|senior secondary)\b"),
]

MONTHS = r"(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE_RANGE = re.compile(
    rf"(?:{MONTHS}\s*)?((?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*(?:(?:{MONTHS}\s*)?((?:19|20)\d{{2}})|(present|current|now|ongoing|till date))",
    re.IGNORECASE,
)
YEARS_STATED = re.compile(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:experience|exp)", re.IGNORECASE)


class ResumeError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        return _pdf_text(data)
    if name.endswith(".docx"):
        return _docx_text(data)
    if name.endswith((".txt", ".md")) or not name:
        return data.decode("utf-8", errors="ignore")
    raise ResumeError("Unsupported file type. Please upload a PDF, DOCX or TXT resume.")


def _pdf_text(data: bytes) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:  # corrupted or encrypted file
            raise ResumeError(f"Could not read the PDF: {exc}") from exc
    if not text.strip():
        raise ResumeError("No text found in the PDF. Scanned image resumes are not supported; please upload a text-based PDF.")
    return text


def _docx_text(data: bytes) -> str:
    try:
        import docx

        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeError(f"Could not read the DOCX file: {exc}") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def split_sections(text: str) -> dict[str, str]:
    """Group resume lines under the nearest recognised heading."""
    lookup = {alias: key for key, aliases in SECTION_HEADINGS.items() for alias in aliases}
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = re.sub(r"[^a-z' ]", "", line.lower()).strip()
        if len(line) <= 40 and heading in lookup:
            current = lookup[heading]
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return {k: "\n".join(v) for k, v in sections.items() if v}


def detect_education(text: str) -> str | None:
    lowered = text.lower()
    for level, pattern in EDUCATION_LEVELS:
        if re.search(pattern, lowered):
            return level
    return None


def estimate_experience_years(experience_text: str, full_text: str) -> float:
    """Use an explicit "N years of experience" if stated, otherwise sum date ranges in the experience section."""
    stated = [float(m.group(1)) for m in YEARS_STATED.finditer(full_text)]
    if stated:
        return min(max(stated), 40.0)
    total = 0.0
    this_year = date.today().year
    for m in DATE_RANGE.finditer(experience_text or ""):
        start = int(m.group(2))
        end = this_year if m.group(5) else int(m.group(4))
        if start <= end <= this_year + 1:
            total += max(end - start, 0.25)  # short internships still count a little
    return round(min(total, 40.0), 1)


def extract_projects(projects_text: str, limit: int = 5) -> list[str]:
    """Pick project titles: short lines, or the part before a separator such as ':' or '|'."""
    projects = []
    for line in (projects_text or "").splitlines():
        line = line.strip(" •●▪-–*\t")
        if not line:
            continue
        title = re.split(r"\s[|:–—-]\s|:\s", line, maxsplit=1)[0].strip()
        if 3 <= len(title) <= 80 and not title.lower().startswith(("tech", "tools", "technologies", "github", "link")):
            words = title.split()
            # Titles are short; long sentences are descriptions of the previous project.
            if len(words) <= 10:
                projects.append(title)
        if len(projects) >= limit:
            break
    return projects


def parse_resume(text: str) -> dict:
    sections = split_sections(text)
    experience_text = sections.get("experience", "")
    return {
        "skills": extract_skills(text),
        "education": detect_education(sections.get("education", "") or text),
        "experience_years": estimate_experience_years(experience_text, text),
        "projects": extract_projects(sections.get("projects", "")),
        "sections_found": [k for k in sections if k != "header"],
        "word_count": len(text.split()),
        "text": text,
    }
