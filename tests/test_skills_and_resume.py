import io

import docx

from app.resume_parser import estimate_experience_years, extract_text, parse_resume
from app.skill_extractor import extract_skills, normalize_skills


def test_extracts_common_skills():
    found = extract_skills("Worked with Python, React.js, Node.js, C++ and AWS (EC2, S3). Familiar with k8s and ML.")
    for skill in ["Python", "React", "Node.js", "C++", "AWS", "Kubernetes", "Machine Learning"]:
        assert skill in found


def test_avoids_substring_false_positives():
    found = extract_skills("JavaScript developer. The rest of the team. 5th sem. TF-IDF features. R&D lab. A. C. Sharma.")
    assert "Java" not in found
    assert "REST APIs" not in found
    assert "TensorFlow" not in found
    assert "R" not in found
    assert "C" not in found


def test_case_sensitive_short_skills():
    assert {"C", "R", "Go"} <= set(extract_skills("Languages: C, R, Go, Python"))


def test_normalize_skills_splits_unknown():
    known, unknown = normalize_skills(["reactjs", "ML", "Power bi", "basket weaving", "react"])
    assert known == ["React", "Machine Learning", "Power BI"]
    assert unknown == ["basket weaving"]


def test_parse_sample_resume(resume_text):
    r = parse_resume(resume_text)
    assert r["education"] == "Bachelors"
    assert {"Python", "SQL", "Machine Learning", "Pandas", "Docker"} <= set(r["skills"])
    assert r["projects"][:3] == ["Crop Yield Predictor", "Resume Screener", "Campus Event App"]
    assert {"education", "experience", "projects", "skills"} <= set(r["sections_found"])
    assert 0 < r["experience_years"] < 1


def test_experience_from_stated_years_and_ranges():
    assert estimate_experience_years("", "Engineer with 4+ years of experience") == 4
    assert estimate_experience_years("Developer, Acme, Jan 2018 - Dec 2021", "") == 3


def test_extract_text_from_docx():
    d = docx.Document()
    d.add_paragraph("SKILLS")
    d.add_paragraph("Python, Django, PostgreSQL")
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text("cv.docx", buf.getvalue())
    assert {"Python", "Django", "PostgreSQL"} <= set(extract_skills(text))


def test_extract_text_from_real_pdf(resume_text):
    from pathlib import Path

    pdf = Path(__file__).resolve().parent.parent / "samples" / "sample_resume.pdf"
    parsed = parse_resume(extract_text("resume.pdf", pdf.read_bytes()))
    assert parsed["skills"] == parse_resume(resume_text)["skills"]
    assert parsed["projects"][:3] == ["Crop Yield Predictor", "Resume Screener", "Campus Event App"]


def test_pdf_upload_via_api(client):
    from pathlib import Path

    pdf = Path(__file__).resolve().parent.parent / "samples" / "sample_resume.pdf"
    r = client.post("/api/resume/parse", files={"file": ("resume.pdf", pdf.read_bytes(), "application/pdf")})
    assert r.status_code == 200 and r.json()["education"] == "Bachelors"
    bad = client.post("/api/resume/parse", files={"file": ("broken.pdf", b"%PDF-1.4 not really a pdf", "application/pdf")})
    assert bad.status_code == 422
